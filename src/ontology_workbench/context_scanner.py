from __future__ import annotations

import csv
import html
import io
import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from dotenv import dotenv_values, load_dotenv

from ontology_workbench.models import DocumentRecord, OntologyProject, utc_now_iso


ROOT_DIR = Path(__file__).resolve().parents[2]
PROMPTS_DIR = ROOT_DIR / "prompts"
DEFAULT_TOOL2_PROMPT_PATH = PROMPTS_DIR / "tool2_business_context_extraction.md"
CONTEXT_CHUNK_SIZE = 6_000
CONTEXT_CHUNK_OVERLAP = 500
LLM_BATCH_CHAR_LIMIT = 18_000

load_dotenv(ROOT_DIR / ".env", override=False)

SUPPORTED_DOCUMENT_TYPES = [
    "functional_docs",
    "technical_docs",
    "kpi_definitions",
    "data_dictionary",
    "process_docs",
    "architecture_docs",
    "bi_exports",
    "misc",
]

SUPPORTED_FILE_EXTENSIONS = [
    "pdf",
    "docx",
    "txt",
    "md",
    "markdown",
    "csv",
    "json",
    "yaml",
    "yml",
    "html",
    "htm",
    "sql",
    "xlsx",
]


@dataclass(slots=True)
class LlmSettings:
    provider: str
    model: str
    api_key_present: bool
    base_url: str = ""
    endpoint: str = ""
    deployment: str = ""
    api_version: str = ""
    configuration_source: str = "local_env"
    reasoning_effort: str = ""
    data_policy: str = "approved_external"
    api_key: str = field(default="", repr=False)

    @property
    def enabled(self) -> bool:
        if self.provider == "disabled":
            return False
        if self.data_policy == "local_only" and self.provider != "ollama":
            return False
        if self.provider == "ollama":
            return True
        return self.api_key_present and bool(self.model)


def load_llm_settings() -> LlmSettings:
    shared_config = _shared_llm_config()
    explicit_provider = _first_env("ONTO_LLM_PROVIDER") or _first_config(
        shared_config, "ONTO_LLM_PROVIDER", "LLM_PROVIDER"
    )
    provider = (explicit_provider or "disabled").lower()
    model = (
        _first_env("ONTO_LLM_MODEL")
        or _first_config(shared_config, "ONTO_LLM_MODEL", "LLM_MODEL")
        or _first_env("AZURE_OPENAI_MODEL", "AZURE_OPENAI_DEPLOYMENT_ID")
    )
    api_key = (
        _first_env("ONTO_LLM_API_KEY")
        or _first_config(shared_config, "ONTO_LLM_API_KEY", "LLM_API_KEY", "OPENAI_API_KEY")
        or _first_env("AZURE_OPENAI_API_KEY")
    )

    if not explicit_provider and provider == "disabled" and _first_env("AZURE_OPENAI_ENDPOINT", "ONTO_AZURE_OPENAI_ENDPOINT"):
        provider = "azure_openai"
    data_policy = (
        _first_env("ONTO_LLM_DATA_POLICY")
        or _first_config(shared_config, "ONTO_LLM_DATA_POLICY", "LLM_DATA_POLICY")
        or "approved_external"
    ).lower()
    if data_policy not in {"approved_external", "local_only"}:
        raise ValueError("ONTO_LLM_DATA_POLICY debe ser approved_external o local_only")

    return LlmSettings(
        provider=provider,
        model=model,
        api_key_present=bool(api_key),
        base_url=(
            _first_env("ONTO_LLM_BASE_URL")
            or _first_config(shared_config, "ONTO_LLM_BASE_URL", "LLM_BASE_URL", "OPENAI_BASE_URL")
            or _first_env("OPENAI_BASE_URL")
        ),
        endpoint=_first_env("ONTO_AZURE_OPENAI_ENDPOINT", "AZURE_OPENAI_ENDPOINT"),
        deployment=_first_env("ONTO_AZURE_OPENAI_DEPLOYMENT", "AZURE_OPENAI_DEPLOYMENT_ID"),
        api_version=_first_env("ONTO_AZURE_OPENAI_API_VERSION", "AZURE_OPENAI_API_VERSION", default="2024-10-21"),
        configuration_source="shared_config" if shared_config and explicit_provider else "local_env",
        reasoning_effort=(
            _first_env("ONTO_LLM_REASONING_EFFORT")
            or _first_config(shared_config, "LLM_REASONING_EFFORT")
            or ("none" if provider == "openai" else "")
        ),
        data_policy=data_policy,
        api_key=api_key,
    )


def extract_text_from_bytes(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".sql", ".yaml", ".yml"}:
        return _decode_text(content)
    if suffix == ".json":
        parsed = json.loads(_decode_text(content))
        return json.dumps(parsed, indent=2, ensure_ascii=False)
    if suffix in {".html", ".htm"}:
        return _html_to_text(_decode_text(content))
    if suffix == ".csv":
        return _csv_to_text(content)
    if suffix == ".pdf":
        return _pdf_to_text(content)
    if suffix == ".docx":
        return _docx_to_text(content)
    if suffix == ".xlsx":
        return _xlsx_to_text(content)
    raise ValueError(f"Formato no soportado para extraccion: {suffix or filename}")


def build_document_chunks(
    documents: list[DocumentRecord],
    document_texts: dict[str, str],
    chunk_size: int = CONTEXT_CHUNK_SIZE,
    overlap: int = CONTEXT_CHUNK_OVERLAP,
) -> list[dict[str, object]]:
    """Split extracted text into stable, locally persisted evidence chunks."""
    chunks: list[dict[str, object]] = []
    for document in documents:
        text = document_texts.get(document.document_id, "").strip()
        if not text:
            continue
        start = 0
        chunk_index = 1
        while start < len(text):
            end = min(len(text), start + chunk_size)
            if end < len(text):
                boundary = text.rfind("\n", start + max(chunk_size // 2, 1), end)
                if boundary > start:
                    end = boundary
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunks.append(
                    {
                        "chunk_id": f"{document.document_id}-chunk-{chunk_index:03d}",
                        "document_id": document.document_id,
                        "filename": document.filename,
                        "chunk_index": chunk_index,
                        "char_start": start,
                        "char_end": end,
                        "text": chunk_text,
                    }
                )
                chunk_index += 1
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return chunks


def build_business_context_inventory(
    project: OntologyProject,
    documents: list[DocumentRecord],
    document_texts: dict[str, str],
    chunks: list[dict[str, object]] | None = None,
) -> dict[str, object]:
    chunks = chunks if chunks is not None else build_document_chunks(documents, document_texts)
    settings = load_llm_settings()
    if settings.enabled:
        try:
            inventory = _build_inventory_with_llm(project, documents, chunks, settings)
            inventory["scan_mode"] = settings.provider
            inventory["chunking"] = _chunking_summary(chunks)
            return inventory
        except Exception as exc:
            fallback = _build_inventory_heuristic(project, documents, chunks)
            fallback["scan_mode"] = "heuristic-fallback"
            fallback["scanner_warning"] = str(exc)
            fallback["chunking"] = _chunking_summary(chunks)
            return fallback

    heuristic = _build_inventory_heuristic(project, documents, chunks)
    heuristic["scan_mode"] = "heuristic"
    heuristic["chunking"] = _chunking_summary(chunks)
    return heuristic


def _build_inventory_with_llm(
    project: OntologyProject,
    documents: list[DocumentRecord],
    chunks: list[dict[str, object]],
    settings: LlmSettings,
) -> dict[str, object]:
    prompt = load_prompt_text(DEFAULT_TOOL2_PROMPT_PATH)
    inventory = _inventory_shell(project, documents)
    keys = _inventory_content_keys()
    for batch in _chunk_batches(chunks):
        context = "\n\n".join(
            (
                f"## Chunk: {chunk['chunk_id']}\n"
                f"Documento: {chunk['filename']}\n"
                f"Rango de caracteres: {chunk['char_start']}-{chunk['char_end']}\n"
                f"Contenido:\n{chunk['text']}"
            )
            for chunk in batch
        )
        messages = [
            {"role": "system", "content": "Eres un analista funcional de datos y negocio."},
            {
                "role": "user",
                "content": (
                    f"Proyecto: {project.name}\n"
                    f"Conceptos tecnicos conocidos: {[concept.name for concept in project.concepts[:80]]}\n\n"
                    f"{prompt}\n\n"
                    f"Extrae solo conocimiento sustentado en los chunks siguientes.\n\n{context}"
                ),
            },
        ]
        parsed = json.loads(_call_llm(messages, settings))
        for key in keys:
            values = parsed.get(key, [])
            if isinstance(values, list):
                inventory[key].extend(values)
    _normalize_inventory_evidence(inventory, chunks)
    return inventory


def _inventory_content_keys() -> tuple[str, ...]:
    return (
        "business_terms",
        "definitions",
        "business_rules",
        "kpis",
        "processes",
        "states",
        "synonyms",
        "candidate_entities",
        "candidate_relationships",
        "ambiguities",
        "questions_for_workshop",
    )


def _chunk_batches(chunks: list[dict[str, object]]) -> list[list[dict[str, object]]]:
    batches: list[list[dict[str, object]]] = []
    batch: list[dict[str, object]] = []
    batch_chars = 0
    for chunk in chunks:
        chunk_chars = len(str(chunk["text"]))
        if batch and batch_chars + chunk_chars > LLM_BATCH_CHAR_LIMIT:
            batches.append(batch)
            batch = []
            batch_chars = 0
        batch.append(chunk)
        batch_chars += chunk_chars
    if batch:
        batches.append(batch)
    return batches


def _normalize_inventory_evidence(
    inventory: dict[str, object], chunks: list[dict[str, object]]
) -> None:
    chunks_by_id = {str(chunk["chunk_id"]): chunk for chunk in chunks}
    for key in _inventory_content_keys():
        values = inventory.get(key, [])
        if not isinstance(values, list):
            inventory[key] = []
            continue
        normalized_values: list[dict[str, object]] = []
        for value in values:
            if not isinstance(value, dict):
                continue
            item = dict(value)
            chunk = chunks_by_id.get(str(item.get("source_chunk_id", "")))
            if chunk:
                item["source_document_id"] = str(chunk["document_id"])
                item["source_chunk_id"] = str(chunk["chunk_id"])
                item["source_excerpt"] = str(item.get("source_excerpt") or str(chunk["text"])[:600])
                item["evidence_status"] = "traceable"
            else:
                item["source_chunk_id"] = ""
                item["source_excerpt"] = ""
                item["evidence_status"] = "missing_chunk_reference"
                if "status" in item:
                    item["status"] = "needs_evidence"
            normalized_values.append(item)
        inventory[key] = normalized_values


def _chunk_evidence(chunk: dict[str, object]) -> dict[str, str]:
    return {
        "source_chunk_id": str(chunk["chunk_id"]),
        "source_excerpt": str(chunk["text"])[:600],
        "evidence_status": "traceable",
    }


def _chunking_summary(chunks: list[dict[str, object]]) -> dict[str, int]:
    return {
        "chunks": len(chunks),
        "chunk_size": CONTEXT_CHUNK_SIZE,
        "overlap": CONTEXT_CHUNK_OVERLAP,
    }


def _call_llm(messages: list[dict[str, str]], settings: LlmSettings) -> str:
    if settings.provider == "openai":
        from openai import OpenAI

        client = OpenAI(
            api_key=settings.api_key,
            base_url=settings.base_url or None,
        )
        response = client.chat.completions.create(
            model=settings.model,
            messages=messages,
            response_format={"type": "json_object"},
            temperature=0.1,
            extra_body={"reasoning_effort": settings.reasoning_effort} if settings.reasoning_effort else None,
        )
        return response.choices[0].message.content or "{}"

    if settings.provider == "azure_openai":
        from openai import AzureOpenAI

        client = AzureOpenAI(
            api_key=settings.api_key,
            api_version=settings.api_version,
            azure_endpoint=settings.endpoint,
        )
        response = client.chat.completions.create(
            model=settings.deployment or settings.model,
            messages=messages,
            response_format={"type": "json_object"},
            temperature=0.1,
        )
        return response.choices[0].message.content or "{}"

    if settings.provider == "ollama":
        import urllib.request

        payload = json.dumps(
            {
                "model": settings.model or "llama3.1:8b",
                "messages": messages,
                "stream": False,
                "format": "json",
            }
        ).encode("utf-8")
        request = urllib.request.Request(
            url=settings.base_url or "http://localhost:11434/api/chat",
            data=payload,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=120) as response:
            body = json.loads(response.read().decode("utf-8"))
        return body.get("message", {}).get("content", "{}")

    raise ValueError(f"Proveedor LLM no soportado: {settings.provider}")


def call_llm_json(messages: list[dict[str, str]], settings: LlmSettings) -> dict[str, object]:
    """Call the configured provider and parse the structured response used by Argos."""
    if not settings.enabled:
        raise ValueError("El proveedor LLM no está habilitado")
    parsed = json.loads(_call_llm(messages, settings))
    if not isinstance(parsed, dict):
        raise ValueError("El LLM no devolvió un objeto JSON")
    return {str(key): value for key, value in parsed.items()}


def _build_inventory_heuristic(
    project: OntologyProject,
    documents: list[DocumentRecord],
    chunks: list[dict[str, object]],
) -> dict[str, object]:
    inventory = _inventory_shell(project, documents)
    seen_terms: set[str] = set()
    seen_rules: set[str] = set()
    seen_kpis: set[str] = set()
    seen_questions: set[str] = set()
    semantic_names = [concept.name.lower() for concept in project.concepts]

    for chunk in chunks:
        document_id = str(chunk["document_id"])
        text = str(chunk["text"])
        evidence = _chunk_evidence(chunk)
        for raw_line in text.splitlines():
            line = raw_line.strip()
            if len(line) < 4:
                continue

            table_cells = _markdown_table_cells(line)
            if table_cells:
                _extract_markdown_table_candidate(
                    inventory, table_cells, document_id, evidence, seen_terms, seen_rules, seen_kpis
                )
                continue

            if ":" in line and len(line.split(":", 1)[0]) <= 80:
                term, definition = [_clean_markdown_text(part) for part in line.split(":", 1)]
                normalized = term.lower()
                if (
                    term
                    and definition
                    and not _is_document_metadata_term(term)
                    and normalized not in seen_terms
                ):
                    seen_terms.add(normalized)
                    inventory["business_terms"].append(
                        {
                            "term": term,
                            "source_document_id": document_id,
                            "confidence": 0.55,
                            "status": "pending_review",
                            **evidence,
                        }
                    )
                    inventory["definitions"].append(
                        {
                            "term": term,
                            "definition": definition,
                            "source_document_id": document_id,
                            "confidence": 0.55,
                            "status": "pending_review",
                            **evidence,
                        }
                    )

            lowered = line.lower()
            cleaned_line = _clean_markdown_text(line)
            if any(keyword in lowered for keyword in ["kpi", "indicador", "métrica", "metrica"]):
                normalized = cleaned_line.lower()
                if normalized not in seen_kpis:
                    seen_kpis.add(normalized)
                    inventory["kpis"].append(
                        {
                            "text": cleaned_line,
                            "source_document_id": document_id,
                            "confidence": 0.5,
                            "status": "pending_review",
                            **evidence,
                        }
                    )

            if any(keyword in lowered for keyword in ["debe", "no debe", "solo si", "excepto"]):
                normalized = cleaned_line.lower()
                if normalized not in seen_rules:
                    seen_rules.add(normalized)
                    inventory["business_rules"].append(
                        {
                            "text": cleaned_line,
                            "source_document_id": document_id,
                            "confidence": 0.5,
                            "status": "pending_review",
                            **evidence,
                        }
                    )

            if "?" in line and line not in seen_questions:
                seen_questions.add(line)
                inventory["questions_for_workshop"].append(
                    {
                        "question": cleaned_line,
                        "source_document_id": document_id,
                        "priority": "medium",
                        **evidence,
                    }
                )

            if any(token in lowered for token in ["proceso", "flujo", "paso"]):
                inventory["processes"].append(
                    {
                        "text": cleaned_line,
                        "source_document_id": document_id,
                        "confidence": 0.45,
                        "status": "pending_review",
                        **evidence,
                    }
                )

        lowered_text = text.lower()
        for concept_name in semantic_names:
            if concept_name and concept_name in lowered_text:
                inventory["candidate_entities"].append(
                    {
                        "term": concept_name,
                        "matched_semantic_object": concept_name,
                        "source_document_id": document_id,
                        "confidence": 0.7,
                        "status": "suggested",
                        **evidence,
                    }
                )

    return inventory


def _markdown_table_cells(line: str) -> list[str]:
    """Return meaningful cells for a Markdown table row, ignoring table decoration."""
    if "|" not in line:
        return []
    cells = [_clean_markdown_text(cell) for cell in line.strip().strip("|").split("|")]
    if len(cells) < 2 or not any(cells):
        return []
    if all(not cell or re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
        return []
    return cells


def _extract_markdown_table_candidate(
    inventory: dict[str, object],
    cells: list[str],
    document_id: str,
    evidence: dict[str, str],
    seen_terms: set[str],
    seen_rules: set[str],
    seen_kpis: set[str],
) -> None:
    """Extract stable facts from common glossary and rule-catalog tables."""
    first = cells[0]
    if not first or _is_markdown_table_header(cells):
        return

    if first.lower().startswith("gold_sic.") and len(cells) >= 2:
        normalized = first.lower()
        if normalized not in seen_terms:
            seen_terms.add(normalized)
            base = {
                "term": first,
                "definition": " ".join(cell for cell in cells[1:] if cell),
                "source_document_id": document_id,
                "confidence": 0.85,
                "status": "pending_review",
                **evidence,
            }
            inventory["business_terms"].append({key: value for key, value in base.items() if key != "definition"})
            inventory["definitions"].append(base)
        return

    if len(cells) == 2 and len(first) <= 80 and len(cells[1]) >= 4:
        normalized = first.lower()
        if normalized not in seen_terms:
            seen_terms.add(normalized)
            base = {
                "term": first,
                "definition": cells[1],
                "source_document_id": document_id,
                "confidence": 0.8,
                "status": "pending_review",
                **evidence,
            }
            inventory["business_terms"].append({key: value for key, value in base.items() if key != "definition"})
            inventory["definitions"].append(base)
        return

    if re.fullmatch(r"\d{3,4}", first) and len(cells) >= 3:
        title = cells[1] or f"Regla {first}"
        rule_text = f"Regla {first} — {title}: " + "; ".join(cell for cell in cells[2:] if cell)
        normalized = rule_text.lower()
        if normalized not in seen_rules:
            seen_rules.add(normalized)
            inventory["business_rules"].append(
                {
                    "text": rule_text,
                    "source_document_id": document_id,
                    "confidence": 0.75,
                    "status": "pending_review",
                    **evidence,
                }
            )
        if any(token in normalized for token in ["%", "porcentaje", "indicador", "metrica", "métrica"]):
            if normalized not in seen_kpis:
                seen_kpis.add(normalized)
                inventory["kpis"].append(
                    {
                        "text": rule_text,
                        "source_document_id": document_id,
                        "confidence": 0.7,
                        "status": "pending_review",
                        **evidence,
                    }
                )


def _clean_markdown_text(value: str) -> str:
    text = value.strip()
    text = re.sub(r"^(?:#{1,6}|>|[-*])\s+", "", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    return re.sub(r"\s+", " ", text).strip()


def _is_document_metadata_term(term: str) -> bool:
    normalized = term.lower().strip()
    return normalized in {"ultima actualizacion", "última actualización", "estado", "objetivo", "alcance"} or normalized.startswith("actualizacion ")


def _is_markdown_table_header(cells: list[str]) -> bool:
    first = cells[0].lower().strip()
    return first in {
        "dato / metadata",
        "dato",
        "codigo",
        "código",
        "regla",
        "descripcion",
        "descripción",
        "nombre",
    }


def _inventory_shell(project: OntologyProject, documents: list[DocumentRecord]) -> dict[str, object]:
    return {
        "client": project.metadata.get("client_id", project.name),
        "domain": project.metadata.get("domain_id", project.id),
        "project_id": project.id,
        "run_id": utc_now_iso(),
        "source_documents": [document.to_dict() for document in documents],
        "business_terms": [],
        "definitions": [],
        "business_rules": [],
        "kpis": [],
        "processes": [],
        "states": [],
        "synonyms": [],
        "candidate_entities": [],
        "candidate_relationships": [],
        "ambiguities": [],
        "questions_for_workshop": [],
    }


def load_prompt_text(prompt_path: Path) -> str:
    if prompt_path.exists():
        return prompt_path.read_text(encoding="utf-8").strip()
    raise FileNotFoundError(f"Prompt file not found: {prompt_path}")


def _first_env(*names: str, default: str = "") -> str:
    for name in names:
        value = os.getenv(name, "").strip()
        if value:
            return value
    return default


def _shared_llm_config() -> dict[str, str]:
    """Read an explicitly configured external LLM file without copying its secrets."""
    config_path = _first_env("ONTO_LLM_CONFIG_PATH")
    if not config_path:
        return {}
    path = Path(config_path).expanduser()
    if not path.is_file():
        return {}
    return {
        str(key): str(value).strip()
        for key, value in dotenv_values(path).items()
        if value is not None and str(value).strip()
    }


def _first_config(config: dict[str, str], *names: str) -> str:
    for name in names:
        value = config.get(name, "").strip()
        if value:
            return value
    return ""


def _decode_text(content: bytes) -> str:
    for encoding in ["utf-8", "utf-8-sig", "latin-1"]:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _csv_to_text(content: bytes) -> str:
    reader = csv.reader(io.StringIO(_decode_text(content)))
    rows = [" | ".join(cell.strip() for cell in row) for row in reader]
    return "\n".join(rows)


def _html_to_text(raw_html: str) -> str:
    without_tags = re.sub(r"<[^>]+>", " ", raw_html)
    return html.unescape(re.sub(r"\s+", " ", without_tags)).strip()


def _pdf_to_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _docx_to_text(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def _xlsx_to_text(content: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        chunks.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                chunks.append(" | ".join(values))
    return "\n".join(chunks)
